#!/usr/bin/env python3
"""Scientific support for each critical limit. English for the reviewer, Spanish to read."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

GREEN = "1F4D36"
INK = RGBColor(0x1E, 0x1A, 0x16)
MUTED = RGBColor(0x4A, 0x43, 0x3A)


def shade(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def p(doc, text, size=11, bold=False, before=0, after=6, center=False, italic=False, color=INK):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(text)
    font(run, size=size, bold=bold, color=color, italic=italic)
    return par


def h1(doc, text):
    return p(doc, text, size=16, bold=True, before=14, after=6, color=RGBColor(0x1F, 0x4D, 0x36))


def h2(doc, text):
    return p(doc, text, size=13, bold=True, before=10, after=4, color=RGBColor(0x1F, 0x4D, 0x36))


def body(doc, text):
    return p(doc, text)


def bullet(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.25)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.space_before = Pt(0)
    run = par.add_run("• " + text)
    font(run, size=11)
    return par


def table(doc, headers, rows, size=8):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        font(run, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, GREEN)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=size)
            if r % 2 == 1:
                shade(cell, "F4F1EA")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def line(doc, label):
    body(doc, label + " " + "_" * 42)


def header_footer(doc, footer):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("Scientific support for critical limits  ·  Carne Seca Jesus Canales, LLC")
    font(r, size=8, color=MUTED)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(footer + "  ·  Page ")
    font(r, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run = fp.add_run()
        font(run, size=8, color=MUTED)
        run._r.append(el)


def english():
    doc = Document()
    header_footer(doc, "Criteria for each critical limit")

    p(doc, "SCIENTIFIC SUPPORT", size=20, bold=True, center=True, after=2, color=RGBColor(0x1F, 0x4D, 0x36))
    p(doc, "Criteria used to set each critical limit", size=14, bold=True, center=True, after=2)
    p(doc, "Ready-to-eat beef jerky (carne seca), heat-treated, shelf-stable", size=12, center=True, after=6)
    p(doc, "CARNE SECA JESUS CANALES, LLC", size=12, bold=True, center=True, after=2)
    p(doc, "411 E Main St, Delta, UT 84624", size=11, center=True, after=2)
    p(doc, "Person in charge: Jesús Canales", size=11, center=True, after=8)

    body(
        doc,
        "This document is the scientific support requested for the critical control points. For each critical limit it names the hazard, the exact criterion used to choose that number, and the page or line in the attached document where that criterion is written. It is written so a reviewer can see why each limit is the limit for this ready-to-eat product.",
    )

    h1(doc, "1. The product these limits apply to")
    table(
        doc,
        ["Item", "This product"],
        [
            ["Product", "Beef jerky (carne seca). The customer does not cook it."],
            ["Beef", "Yes. Inspected whole-muscle beef."],
            ["Process", "Heat-treated, then dried. Shelf-stable. Not refrigerated for safety."],
            ["Category", "Heat-treated — shelf stable, 9 CFR 417.2(b)(1)."],
            ["Not this product", "Not fermented. Not acidified. No nitrite, no brine, no starter culture. Nothing is added after the heat step. pH is not a critical limit."],
            ["Order of the steps", "Cold raw beef, then a humid cook, then drying. The meat is not dried before the cook."],
        ],
        size=9,
    )

    h1(doc, "2. Documents that identify the criteria")
    body(doc, "Three documents are attached with this support. Each limit below points to one of them.")
    table(
        doc,
        ["Document", "What it is used for"],
        [
            [
                "HACCP Field Verification Checklist — Jerky (Fully Cooked, Shelf Stable, RTE Meat and Poultry)",
                "Prints the three critical limits: cooler < 5°C (41°F); dry bulb > 77°C (170°F) in 30 minutes and wet bulb 52°C (125°F) to 61°C (142°F); water activity < 0.88. The heat-lethality line is marked “Per FSIS Performance Standard.”",
            ],
            [
                "FSIS Cooking Guideline for Meat and Poultry Products (Revised Appendix A), December 2021, FSIS-GD-2021-14",
                "This is the performance standard named on the jerky checklist. The checklist footnote cites the 1999 web address. The copy attached is the December 2021 revision of that guideline. It supplies the internal endpoint, the come-up time, and the requirement that humidity be applied during lethality, before drying.",
            ],
            [
                "Principles of Preservation of Shelf-Stable Dried Meat Products, FSRE, October 31, 2011",
                "Identifies water activity, not the moisture-to-protein ratio, as the safety measurement for dried meat. States that beef jerky generally has a water activity below 0.88. States that a beef product must address a 5-log reduction of E. coli O157:H7 and that heating is how that reduction is achieved. Lists the minimum water activity for growth of the relevant bacteria.",
            ],
        ],
        size=8,
    )
    body(
        doc,
        "9 CFR 417.5(a)(2) is the requirement to keep the support for the critical limits. 9 CFR 417.4(a)(1) has two parts. This document is the scientific-support part: why the numbers were chosen. The in-plant part is the lot logs that show this oven, this load, and this thickness met those numbers. The validation date on the HACCP plan stays blank until those logs have been reviewed and signed. This document does not cite a laboratory challenge study of this jerky.",
    )

    h1(doc, "3. One-page crosswalk")
    table(
        doc,
        ["CCP", "Critical limit in the plan", "Criterion used to set it", "Where it is written"],
        [
            [
                "CCP-1 Cooler",
                "Product temperature < 41°F (5°C)",
                "Keep raw beef cold so Staphylococcus aureus does not grow and form heat-stable enterotoxin before the cook.",
                "Jerky checklist, Cooler Storage. Appendix A, page 14, heat-stable enterotoxin.",
            ],
            [
                "CCP-2 Dry bulb",
                "> 170°F (77°C) within 30 minutes",
                "The oven is hot on the schedule printed on the jerky checklist. This number does not replace the internal temperature.",
                "Jerky checklist, Heat Lethality.",
            ],
            [
                "CCP-2 Wet bulb",
                "125°F to 142°F (52°C to 61°C) during the lethality stage",
                "Humidity is present during the cook, before drying, and is recorded as wet bulb. The wet-bulb range alone is not Appendix A Humidity Option 1, 2, 3, or 4.",
                "Jerky checklist, Heat Lethality. Appendix A, pages 25–27.",
            ],
            [
                "CCP-2 Internal",
                "≥ 158°F (70°C) in the thickest piece, coldest spot. Not an average.",
                "Table 2: the required log reduction is reached at once when every part of the meat is at 158°F or above, if humidity and come-up time are also met.",
                "Appendix A, Table 2, page 35, footnote 5.",
            ],
            [
                "CCP-2 Come-up",
                "Time between 50°F and 130°F internal ≤ 6 hours",
                "Come-up option that holds S. aureus growth to ≤ 2-log and prevents enterotoxin formation.",
                "Appendix A, pages 23–24 and Table 2 footnote 7.",
            ],
            [
                "CCP-3 Drying",
                "Water activity < 0.88 on the finished jerky, before packing. Highest of three pieces.",
                "Checklist limit for drying. Water activity, not moisture-to-protein ratio, is the safety measurement. Beef jerky is described at a water activity below 0.88.",
                "Jerky checklist, Drying. Preservation document, water-activity section and Jerky Products section.",
            ],
        ],
        size=7,
    )

    h1(doc, "4. CCP-1 — Cooler storage")
    h2(doc, "Limit")
    body(doc, "Product temperature below 5°C (41°F). The probe is in the meat, not in the cooler air. The reading is taken at receiving, at least once a shift while beef for this jerky is in the cooler, and at the start and end of marination.")
    h2(doc, "Hazard")
    body(doc, "Growth of bacteria in the raw beef before the cook, in particular Staphylococcus aureus. Appendix A, page 14, states that S. aureus causes illness when it grows to high levels and produces one or more heat-stable enterotoxins (Kadariya et al., 2014, as cited there). A toxin formed in the raw meat is not destroyed by the heat step that follows. The cook is not a corrective action for warm raw beef.")
    h2(doc, "Criterion")
    body(doc, "The jerky checklist, in the critical-control-point table, line “Cooler Storage,” identifies the critical limit as temperature < 5°C (41°F). That printed limit is the criterion used. Holding the product below 41°F also keeps it below 50°F, which is the bottom of the come-up range in Appendix A. The come-up clock in CCP-2 does not start while the meat is held under this cooler limit.")
    h2(doc, "If the limit is missed")
    body(doc, "The beef is held. Beef at 41°F or above is not made into this jerky. Jesús Canales is told, and the corrective-action form is filled in.")

    h1(doc, "5. CCP-2 — Heat lethality, before drying")
    body(doc, "This is one critical control point with four numbers. All four are required. Drying does not start until this step has passed. The hazards this step is designed to destroy are Salmonella and, because the product contains beef, Shiga toxin-producing E. coli, including E. coli O157:H7. Appendix A, page 13, lists those hazards and states that Salmonella is the indicator of lethality because its destruction indicates destruction of most other pathogens (64 FR 732).")

    h2(doc, "5.1 Dry bulb above 170°F within 30 minutes")
    body(doc, "Criterion: the jerky checklist, line “Heat Lethality,” prints “> 77°C (170°F) dry bulb in 30 min.” That is the criterion for this number. It is an oven condition at the start of lethality. It is not the internal endpoint, and it is not used by itself to claim the log reduction.")

    h2(doc, "5.2 Wet bulb from 125°F to 142°F during the lethality stage")
    body(doc, "Criterion from the checklist: the same Heat Lethality line prints “52°C (125°F) – 61°C (142°F) wet bulb,” with the footnote “Per FSIS Performance Standard.”")
    body(doc, "Criterion from Appendix A, which is that performance standard:")
    bullet(doc, "Page 25: the time-temperature tables use relative humidity as a critical operating parameter so the cook stays moist and the surface is lethal. An establishment that uses those tables must address humidity.")
    bullet(doc, "Page 26: humidity needs to be applied during the lethality treatment, before drying. Drying first, and then a moist cook, is described as a vulnerability. Humidity is monitored through the lethality treatment with wet-bulb and dry-bulb thermometers, or a humidity sensor, for every lot.")
    bullet(doc, "Page 26, Table 1: the four FSIS humidity options are continuous steam, a sealed oven, 90 percent relative humidity for 25 percent of the cook or 1 hour (whichever is longer), or 90 percent relative humidity for the entire cook.")
    bullet(doc, "Page 27: the 2014 jerky guideline’s suggestion of a wet bulb of 125–130°F for 1 hour (about 27–32 percent relative humidity) is not, by itself, adequate to show that the process matches those humidity options. All of the critical operating parameters in the guideline have to be met.")
    body(doc, "How this plan uses those two criteria together: the wet-bulb range on the checklist is the number the monitor writes down, so humidity during the cook is measured and not assumed. This plan does not claim that the wet-bulb range alone is Humidity Option 1, 2, 3, or 4. The lethality claim uses the wet-bulb record together with the internal temperature and the come-up time, and the cook happens before any drying. If the oven later follows one named option in Table 1 (sealed oven or continuous steam), that option will be written on the heat log as the method, with the records Appendix A describes for that option. Until that method is written, the support for humidity is the checklist wet-bulb limit, monitored every lot, during lethality, before drying.")

    h2(doc, "5.3 Internal temperature at or above 158°F")
    body(doc, "Criterion: Appendix A, Table 2, page 35. The table states that the temperatures are the minimum internal temperatures that must be met in all parts of the meat product. Footnote 5 states that the required log reduction is achieved instantly (0 seconds) when the internal temperature of a cooked meat product reaches 158°F or above. Footnote 6 states that relative humidity is a critical operating parameter when the table is used. Footnote 7 states the come-up recommendation of 6 hours or less between 50°F and 130°F.")
    body(doc, "The probe goes into the thickest slice, in the coldest part of the oven. Readings are not averaged. The lowest reading is the lot result. A surface reading is not used for this endpoint. Appendix A, page 24, states that a surface temperature is not support for the endpoint.")
    body(doc, "Beef, separate from the Salmonella indicator: Principles of Preservation of Shelf-Stable Dried Meat Products, page 12, states that processors must validate a 5-log kill of E. coli O157:H7 for products containing beef, and that for the most part the product must be heated to achieve that reduction. This heat step is the step used for that purpose. Water activity is not offered as a substitute for the heat step. Page 2 of that document states that if pathogens are still viable, the product is adulterated, and that inactivation (heat) is part of the safety of dried meat along with the low water activity.")

    h2(doc, "5.4 Come-up time of 6 hours or less")
    body(doc, "Criterion: Appendix A, pages 23–24. The come-up-time option is: total time the product temperature is between 50°F and 130°F is 6 hours or less. The guideline states that this option supports control of S. aureus growth, specifically ≤ 2-log, and that it also prevents enterotoxin formation. The temperatures in that option are internal temperatures. Table 2, footnote 7, repeats the same limit as a critical operating parameter for the table.")
    body(doc, "The clock is the total time in that range. It does not start over if the first heating missed the lethal temperature. Appendix A, page 25, states that a new 6-hour come-up applies to a second cook only after a lethal time-temperature was already achieved.")

    h2(doc, "If any CCP-2 number is missed")
    body(doc, "The lot is held. Drying does not start. The lot is not shipped as ready-to-eat jerky. Heating may continue only when the wet bulb has stayed inside 125–142°F and the only miss is that the internal temperature is not at 158°F yet, and the come-up has not passed 6 hours. A pathogen test on the finished piece is not the release step.")

    h1(doc, "6. CCP-3 — Drying, after CCP-2 has passed")
    h2(doc, "Limit")
    body(doc, "Water activity below 0.88. Three pieces from different places in the oven, including the thickest piece, are read on a calibrated meter after drying and before packing. The highest reading is the lot result. If that reading is 0.88 or above, the lot is not packed.")
    h2(doc, "Hazard")
    body(doc, "Growth of pathogens in a shelf-stable bag if too much water remains available. The product is not refrigerated for safety after packing.")
    h2(doc, "Criteria")
    bullet(doc, "The jerky checklist, line “Drying,” prints the critical limit as water activity < 0.88. That printed limit is the critical limit in the plan.")
    bullet(doc, "Principles of Preservation of Shelf-Stable Dried Meat Products, water-activity section: water activity is the most important single factor for shelf stability of most dried meats. Dried hams, coppa, and beef jerky generally have a water activity less than 0.88. That sentence is the product-specific criterion for using 0.88 as the jerky number.")
    bullet(doc, "The same document, Jerky Products section: FSIS jerky guidelines rely on water activity, not on the moisture-to-protein ratio, as the indicator of final-product safety and shelf stability, and they call for a humidity step at the beginning of the process. The moisture-to-protein ratio is a labeling description. It is not the critical limit. The document gives beef jerky an example ratio of 0.75:1 and states that such ratios are not necessarily indicative of microbial safety.")
    bullet(doc, "The same document lists minimum water activity for growth when other conditions are optimal: E. coli O157:H7, 0.95; Salmonella, 0.94; Listeria monocytogenes, 0.92; S. aureus with air, 0.85. A later sentence states that S. aureus can grow at a water activity as low as 0.86 when oxygen is present. A finished water activity below 0.88 is below the listed minima for E. coli O157:H7, Salmonella, and L. monocytogenes. It is not below 0.85.")
    body(doc, "What this plan does with that last point: CCP-3 is not offered as the step that stops S. aureus from making toxin. Toxin is controlled earlier. CCP-1 keeps the raw beef below 41°F. CCP-2 limits the time between 50°F and 130°F to 6 hours, which is the Appendix A option for preventing enterotoxin, and then applies the lethal cook. CCP-3 is the drying limit on the jerky checklist, supported by the statement that beef jerky is at a water activity below 0.88, and it is below the growth minima listed for E. coli O157:H7, Salmonella, and L. monocytogenes.")
    body(doc, "pH is not a critical limit. The preservation document treats pH as a control for fermented dried meats. This jerky is not fermented and is not acidified.")
    h2(doc, "If the limit is missed")
    body(doc, "The jerky is not packed. If CCP-2 passed, it may be dried longer and read again. If the highest reading is still 0.88 or above, it is not shipped. It is not wetted and cooked again under this plan.")

    h1(doc, "7. How the records are examined, without a routine laboratory test")
    body(doc, "The jerky checklist asks whether the records for the day match what is happening, and whether monitoring follows the plan. It does not require a laboratory test of every bag. Before a lot ships, a reviewer writes Met or Not met against the number on each log:")
    bullet(doc, "Product temperature < 41°F.")
    bullet(doc, "Dry bulb > 170°F within 30 minutes.")
    bullet(doc, "Wet bulb between 125°F and 142°F during the lethality stage.")
    bullet(doc, "Come-up from 50°F to 130°F internal, 6 hours or less.")
    bullet(doc, "Internal temperature ≥ 158°F on the lowest slice, not an average.")
    bullet(doc, "Highest water activity < 0.88.")
    bullet(doc, "That day’s calibration was acceptable, and sanitation was signed.")
    body(doc, "A blank, a check mark, or the word “yes” where a number belongs is Not met. The lot stays. That review is the examination of the information. A laboratory test is not the way a failed lot is released.")

    h1(doc, "8. What is still filled in at the plant")
    body(doc, "This support identifies the criteria. It does not replace the plant records. The following stay blank until the plant completes them:")
    bullet(doc, "Phone, e-mail, and establishment number.")
    bullet(doc, "Name and training date of the HACCP-trained person (9 CFR 417.7).")
    bullet(doc, "Signature of Jesús Canales on the HACCP plan.")
    bullet(doc, "Validation date, after logs from this oven, this load, and this slice thickness show that CCP-1, CCP-2, and CCP-3 were met (9 CFR 417.4).")
    body(doc, "A thicker slice, a different oven, a heavier load, a brine, a dry cure, or a spice added after the cook is not supported by this document.")

    h1(doc, "9. Signature")
    body(doc, "I am providing this document as the scientific support for the critical limits in the HACCP plan for this ready-to-eat beef jerky.")
    line(doc, "Jesús Canales:")
    line(doc, "Date:")
    line(doc, "HACCP-trained person:")
    line(doc, "Date:")

    p(
        doc,
        "Attachments: (1) Jerky field verification checklist. (2) FSIS-GD-2021-14, December 2021. (3) Principles of Preservation of Shelf-Stable Dried Meat Products, October 31, 2011. (4) HACCP plan for this product.",
        size=10,
        italic=True,
        before=8,
    )
    doc.save("/workspace/haccp-plan/Scientific-Support-CCPs-Carne-Seca.docx")


def spanish():
    doc = Document()
    header_footer(doc, "Respaldo científico de cada límite")
    # Spanish header text is set in English function's header; override by rewriting header run is already English.
    # Replace header for this document.
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.clear()
    r = hp.add_run("Respaldo científico  ·  Carne Seca Jesus Canales, LLC")
    font(r, size=8, color=MUTED)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p(doc, "RESPALDO CIENTÍFICO", size=20, bold=True, center=True, after=2, color=RGBColor(0x1F, 0x4D, 0x36))
    p(doc, "Por qué cada límite deja la carne seca segura para comer", size=13, bold=True, center=True, after=2)
    p(doc, "Carne seca de res, lista para comer, estable sin refrigeración", size=12, center=True, after=6)
    p(doc, "CARNE SECA JESUS CANALES, LLC", size=12, bold=True, center=True, after=2)
    p(doc, "411 E Main St, Delta, UT 84624", size=11, center=True, after=2)
    p(doc, "Responsable: Jesús Canales", size=11, center=True, after=8)

    body(
        doc,
        "Este es el documento completo que pide el revisor. El archivo en inglés, para enviarlo, se llama Scientific-Support-CCPs-Carne-Seca.docx y está en esta misma carpeta. Este archivo en español dice lo mismo, para que se pueda leer y firmar sabiendo qué se está entregando.",
    )
    body(
        doc,
        "El revisor pidió que el respaldo identifique el criterio usado para fijar cada límite crítico, de modo que el producto listo para comer sea seguro. Abajo, cada límite tiene el peligro, el criterio y el documento donde está escrito ese criterio.",
    )

    h1(doc, "1. De qué producto se habla")
    bullet(doc, "Carne seca de res. El cliente no la cocina.")
    bullet(doc, "Sí lleva res. Es músculo entero inspeccionado.")
    bullet(doc, "Se cocina con humedad y después se seca. No se seca primero.")
    bullet(doc, "No es fermentada. No lleva nitrito, ni salmuera, ni cultivo. No se le agrega nada después de la cocción. El pH no es límite.")
    bullet(doc, "Se vende sin refrigeración. La seguridad no depende de que el cliente la guarde en frío.")

    h1(doc, "2. Los tres documentos que se adjuntan")
    table(
        doc,
        ["Documento", "Para qué sirve"],
        [
            ["Hoja de verificación de jerky (cocida, estable, lista para comer)", "Trae escritos los límites: cámara menor de 41 °F; bulbo seco mayor de 170 °F en 30 minutos; bulbo húmedo de 125 °F a 142 °F; actividad de agua menor de 0,88. La cocción dice “según la norma de FSIS”."],
            ["Appendix A de FSIS, diciembre 2021, FSIS-GD-2021-14", "Es esa norma. La hoja cita la dirección de 1999. El archivo que se adjunta es la revisión de diciembre de 2021. De ahí salen la temperatura interna, las 6 horas de subida y la regla de que la humedad va durante la cocción, antes del secado."],
            ["Principles of Preservation of Shelf-Stable Dried Meat Products, 31 de octubre de 2011", "Dice que la seguridad se mide con actividad de agua, no con la relación humedad/proteína. Dice que la carne seca de res suele estar bajo 0,88. Dice que la res tiene que lograr una reducción de 5 log de E. coli O157:H7 y que eso se logra calentando."],
        ],
        size=8,
    )
    body(doc, "Este documento explica por qué se eligieron los números. No reemplaza los registros del horno de Delta. La fecha de validación del plan se llena cuando esos registros muestren que este horno, esta carga y este grosor cumplieron los límites. No hay un estudio de laboratorio de desafío de esta carne seca, y este documento no dice que lo haya.")

    h1(doc, "3. Tabla corta")
    table(
        doc,
        ["Punto", "Límite", "Criterio"],
        [
            ["CCP-1 Cámara", "Producto menor de 41 °F (5 °C)", "La carne cruda no se calienta, para que Staphylococcus aureus no forme una toxina que el horno no destruye."],
            ["CCP-2 Bulbo seco", "Mayor de 170 °F en 30 minutos", "El horno está caliente en el tiempo que trae la hoja de jerky. No sustituye la temperatura interna."],
            ["CCP-2 Bulbo húmedo", "Entre 125 °F y 142 °F durante la cocción", "Hay humedad durante la cocción, antes de secar, y se anota. Ese rango solo no es una de las cuatro opciones de humedad del Appendix A."],
            ["CCP-2 Interna", "158 °F o más, en la tira más gruesa y en el punto más frío", "La Tabla 2 dice que la reducción se logra al instante al llegar a 158 °F en toda la carne, si también se cumplen la humedad y el tiempo de subida."],
            ["CCP-2 Subida", "Entre 50 °F y 130 °F, 6 horas o menos", "Opción del Appendix A para que S. aureus no pase de unos 2 log y no forme toxina."],
            ["CCP-3 Secado", "Actividad de agua menor de 0,88", "Límite escrito en la hoja. La actividad de agua, no la relación humedad/proteína, es la medida de seguridad. La carne seca de res se describe bajo 0,88."],
        ],
        size=8,
    )

    h1(doc, "4. CCP-1 — Cámara")
    body(doc, "Límite: temperatura del producto menor de 41 °F (5 °C). El termómetro va en la carne, no en el aire de la cámara. Se mide al recibir, al menos una vez por turno mientras haya res para esta carne seca, y al empezar y al terminar el marinado.")
    body(doc, "Peligro: que las bacterias crezcan en la res cruda antes del horno. El que más importa aquí es Staphylococcus aureus. El Appendix A, página 14, dice que enferma cuando crece mucho y forma enterotoxinas que resisten el calor. Si la toxina ya se formó en la carne cruda, la cocción no la quita. Por eso una res que llegó a 41 °F o más no entra a este proceso.")
    body(doc, "Criterio: la hoja de jerky, en la línea Cooler Storage, escribe el límite < 5 °C (41 °F). Ese número impreso es el criterio. Además, 41 °F está debajo de 50 °F, que es donde empieza a contar el tiempo de subida del CCP-2. Mientras la carne cumple la cámara, ese reloj no corre.")

    h1(doc, "5. CCP-2 — Cocción con humedad, antes de secar")
    body(doc, "Es un solo punto crítico con cuatro números. Los cuatro se cumplen. El secado no empieza hasta que este paso pasó. Los peligros que este paso está hecho para destruir son Salmonella y, porque es res, E. coli productora de toxina Shiga, incluida O157:H7. El Appendix A, página 13, dice que Salmonella es el indicador: si el calor la destruye, indica que destruye a la mayoría de los otros patógenos.")

    h2(doc, "5.1 Bulbo seco mayor de 170 °F en 30 minutos")
    body(doc, "La hoja de jerky escribe “> 77 °C (170 °F) dry bulb in 30 min”. Es la condición del horno al empezar. No es la temperatura de adentro de la carne y no se usa solo para decir que ya hubo reducción de bacterias.")

    h2(doc, "5.2 Bulbo húmedo de 125 °F a 142 °F")
    body(doc, "La misma línea de la hoja escribe el bulbo húmedo de 52 °C (125 °F) a 61 °C (142 °F), con la nota “según la norma de FSIS”.")
    body(doc, "Lo que dice esa norma, el Appendix A de diciembre de 2021:")
    bullet(doc, "Página 25: las tablas de tiempo y temperatura usan la humedad como parámetro crítico para que la cocción esté húmeda y la superficie también se cocine.")
    bullet(doc, "Página 26: la humedad va durante la letalidad, antes del secado. Secar primero y cocinar húmedo después es una falla. La humedad se vigila en cada lote con bulbo seco y bulbo húmedo, o con un sensor.")
    bullet(doc, "Página 26, Tabla 1: las cuatro opciones de FSIS son vapor continuo, horno cerrado, 90 % de humedad relativa durante el 25 % de la cocción o 1 hora (lo que sea más largo), o 90 % durante toda la cocción.")
    bullet(doc, "Página 27: la guía de jerky de 2014, que sugería bulbo húmedo de 125–130 °F durante 1 hora, no basta por sí sola para decir que el proceso cumple esas opciones. Hay que cumplir todos los parámetros críticos.")
    body(doc, "Cómo se usan las dos cosas: el rango de la hoja es el número que se anota, para que la humedad quede medida. Este plan no dice que ese rango, solo, sea la opción 1, 2, 3 o 4. La afirmación de letalidad junta el bulbo húmedo, la temperatura interna y las 6 horas, y la cocción ocurre antes de secar. Si más adelante el horno trabaja como horno cerrado o con vapor continuo, eso se escribe en la hoja de cocción, con los registros que el Appendix A pide para esa opción.")

    h2(doc, "5.3 Temperatura interna de 158 °F o más")
    body(doc, "Criterio: Appendix A, Tabla 2, página 35. La tabla dice que la temperatura es la mínima interna en todas las partes de la carne. La nota 5 dice que la reducción exigida se logra al instante (0 segundos) cuando la temperatura interna llega a 158 °F o más. La nota 6 dice que la humedad relativa es parámetro crítico para usar la tabla. La nota 7 repite las 6 horas entre 50 °F y 130 °F.")
    body(doc, "La sonda va en la tira más gruesa, en el lugar más frío del horno. No se promedian lecturas. La lectura más baja es el resultado del lote. La temperatura de la superficie no sirve para este límite.")
    body(doc, "Para la res, aparte de Salmonella: el documento de conservación, página 12, dice que hay que validar una reducción de 5 log de E. coli O157:H7 en productos con res, y que en general eso se logra calentando. Este paso de calor es el que se usa para eso. La actividad de agua no sustituye el calor. Si los patógenos siguen vivos, el documento dice que el producto está adulterado.")

    h2(doc, "5.4 Subida de 6 horas o menos")
    body(doc, "Criterio: Appendix A, páginas 23 y 24. El tiempo total en que la temperatura interna está entre 50 °F y 130 °F es de 6 horas o menos. Esa opción sirve para controlar el crecimiento de S. aureus a unos 2 log o menos, y para evitar que se forme la toxina. Si la primera cocción no llegó a la temperatura letal, el reloj no empieza de nuevo.")

    h2(doc, "Si falta un número del CCP-2")
    body(doc, "El lote se detiene. No se seca y no se vende como carne seca lista para comer. Se puede seguir calentando solo si el bulbo húmedo se mantuvo entre 125 °F y 142 °F, lo único que falta es llegar a 158 °F internos, y todavía no se pasaron las 6 horas. Una prueba de laboratorio del producto terminado no libera el lote.")

    h1(doc, "6. CCP-3 — Secado, después de que el CCP-2 pasó")
    body(doc, "Límite: actividad de agua menor de 0,88. Se miden tres piezas de distintos lugares del horno, incluida la más gruesa, con el medidor calibrado, después de secar y antes de empacar. La lectura más alta es la del lote. Si esa lectura es 0,88 o más, no se empaca.")
    body(doc, "Criterios:")
    bullet(doc, "La hoja de jerky, línea Drying, escribe actividad de agua < 0,88. Ese es el límite del plan.")
    bullet(doc, "El documento de conservación dice que la actividad de agua es el factor más importante para que la carne seca se mantenga estable, y que el jamón seco, la coppa y la carne seca de res suelen estar bajo 0,88.")
    bullet(doc, "El mismo documento, en Jerky Products, dice que las guías de FSIS usan la actividad de agua, no la relación humedad/proteína, como indicador de seguridad, y que la humedad va al principio del proceso. La relación humedad/proteína es de etiqueta. No es el límite. El ejemplo de 0,75:1 para carne seca no demuestra por sí solo que el lote sea seguro.")
    bullet(doc, "El mismo documento lista el mínimo de actividad de agua para que crezcan, cuando todo lo demás les favorece: E. coli O157:H7, 0,95; Salmonella, 0,94; Listeria, 0,92; S. aureus con aire, 0,85. Más adelante dice que S. aureus puede crecer hasta en 0,86 si hay oxígeno. Menos de 0,88 queda debajo de E. coli, Salmonella y Listeria. No queda debajo de 0,85.")
    body(doc, "Qué hace el plan con ese último dato: el secado no es el paso que evita la toxina de S. aureus. Esa toxina se controla antes. La cámara mantiene la res bajo 41 °F. La cocción limita el tiempo entre 50 °F y 130 °F a 6 horas, que es la opción del Appendix A para evitar la toxina, y después aplica el calor letal. El secado es el límite de la hoja, apoyado en que la carne seca de res se describe bajo 0,88, y queda debajo de los mínimos de E. coli O157:H7, Salmonella y Listeria.")
    body(doc, "El pH no es límite. El documento de conservación usa el pH en carnes secas fermentadas. Esta no es fermentada.")

    h1(doc, "7. Cómo se revisa el lote, sin laboratorio de rutina")
    body(doc, "Antes de que el lote salga, alguien anota Cumple o No cumple junto al número medido:")
    bullet(doc, "Producto menor de 41 °F.")
    bullet(doc, "Bulbo seco mayor de 170 °F dentro de 30 minutos.")
    bullet(doc, "Bulbo húmedo entre 125 °F y 142 °F durante la cocción.")
    bullet(doc, "De 50 °F a 130 °F internos, 6 horas o menos.")
    bullet(doc, "Interna de 158 °F o más, la lectura más baja.")
    bullet(doc, "Actividad de agua más alta, menor de 0,88.")
    bullet(doc, "Calibración del día en regla y limpieza firmada.")
    body(doc, "Un espacio en blanco, una paloma o la palabra “sí” donde tiene que ir un número, es No cumple. El lote no sale. Así se examina la información. Una prueba de laboratorio no es la forma de soltar un lote que falló.")

    h1(doc, "8. Lo que todavía se llena en la planta")
    bullet(doc, "Teléfono, correo y número de establecimiento.")
    bullet(doc, "Nombre y fecha de la persona capacitada en HACCP.")
    bullet(doc, "Firma de Jesús Canales en el plan y en este respaldo.")
    bullet(doc, "Fecha de validación, cuando los registros de este horno muestren que los tres puntos se cumplieron.")
    body(doc, "Una tira más gruesa, otro horno, una carga más pesada, una salmuera o una especia agregada después de la cocción no están cubiertos por este documento.")

    h1(doc, "9. Firma")
    body(doc, "Entrego este documento como el respaldo científico de los límites críticos del plan HACCP de esta carne seca lista para comer.")
    line(doc, "Jesús Canales:")
    line(doc, "Fecha:")
    line(doc, "Persona capacitada en HACCP:")
    line(doc, "Fecha:")

    p(
        doc,
        "Se adjuntan: (1) la hoja de verificación de jerky, (2) el Appendix A de diciembre de 2021, (3) el documento de conservación de carne seca del 31 de octubre de 2011, y (4) el plan HACCP.",
        size=10,
        italic=True,
        before=8,
    )
    doc.save("/workspace/haccp-plan/Respaldo-Cientifico-CCPs-Carne-Seca.docx")


if __name__ == "__main__":
    english()
    spanish()
    print("ok")
